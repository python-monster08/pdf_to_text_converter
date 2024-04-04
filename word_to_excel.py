import re
from docx import Document
from openpyxl import Workbook


def parse_document(doc_path):
    doc = Document(doc_path)
    data = []

    for para in doc.paragraphs:
        try:
            # Splitting the paragraph by lines
            parts = para.text.split('\n')
            if len(parts) < 3:
                continue  # Skip if there aren't enough parts

            question_info = parts[0].split(' ')
            question_number = question_info[0].strip('.')
            exam_year = re.search(r'\[(\d+)\]', parts[0]).group(1)
            exam_name = 'U.P .P .C.S.'
            exam_part = 'Mains'
            question_statement = parts[0]

            list_I_title = 'List-I (years)'
            list_II_title = 'List-II (events)'

            # Assuming two spaces as a separator for options
            list_options = parts[1].strip().split('  ')
            list_I_options = [opt.strip() for opt in list_options[:4]]
            list_II_options = [opt.strip() for opt in list_options[4:8]]

            codes = parts[2].split('Codes:')[1].strip()
            options_text = re.search(r'Codes:.*?(Answer -)', parts[2], re.DOTALL).group(0)
            options = [opt.strip() for opt in options_text.split(';') if opt]

            correct_answer = re.search(r'Answer -\((\w)\)', parts[2]).group(1)

            record = [
                question_number, exam_name, exam_part, exam_year, question_statement,
                list_I_title, list_II_title, *list_I_options, *list_II_options, codes, *options[:-1], correct_answer
            ]
            data.append(record)
        except Exception as e:
            print(f"An error occurred while parsing: {e}")
            continue  # Skip this paragraph and move to the next

    return data



def create_excel(data, excel_path):
    wb = Workbook()
    ws = wb.active

    # Define headers based on the provided structure
    headers = [
        'Question Number', 'Exam Name', 'Exam Part', 'Exam Year', 'Question Statement',
        'List I Title', 'List II Title', 'List I Option A', 'List I Option B', 'List I Option C', 'List I Option D',
        'List II Option 1', 'List II Option 2', 'List II Option 3', 'List II Option 4', 'Codes',
        'Option A', 'Option B', 'Option C', 'Option D', 'Correct Answer'
    ]
    ws.append(headers)

    for record in data:
        ws.append(record)

    wb.save(excel_path)


# Replace with the correct path to your Word document
doc_data = parse_document('sample_questions.docx')
create_excel(doc_data, 'output.xlsx')

